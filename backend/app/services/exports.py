import csv
import io
from datetime import datetime
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

FORMATS = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


def build_export(fmt: str, title: str, columns: list[str], rows: list[list], subtitle: str = "") -> tuple[bytes, str, str]:
    """Returns (content, media_type, filename) for a tabular report."""
    stamp = datetime.now().strftime("%Y%m%d-%H%M")
    filename = f"{title.lower().replace(' ', '-')}-{stamp}.{fmt}"
    rows = [["" if c is None else c for c in r] for r in rows]

    if fmt == "csv":
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(columns)
        # Prefix formula-looking cells so spreadsheet apps don't execute them.
        writer.writerows([[("'" + c) if isinstance(c, str) and c[:1] in "=+-@" else c for c in r] for r in rows])
        return ("﻿" + buf.getvalue()).encode("utf-8"), FORMATS[fmt], filename

    if fmt == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = title[:31]
        ws.append(columns)
        for r in rows:
            ws.append([("'" + c) if isinstance(c, str) and c[:1] in "=+-@" else c for c in r])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0F766E")
            cell.alignment = Alignment(vertical="center")
        for i, col in enumerate(columns, start=1):
            width = max([len(str(col))] + [len(str(r[i - 1])) for r in rows[:200]]) + 2
            ws.column_dimensions[get_column_letter(i)].width = min(width, 60)
        ws.freeze_panes = "A2"
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue(), FORMATS[fmt], filename

    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=1.2 * cm, rightMargin=1.2 * cm, topMargin=1.2 * cm, bottomMargin=1.2 * cm)
    styles = getSampleStyleSheet()
    cell_style = styles["BodyText"]
    cell_style.fontSize = 8
    cell_style.leading = 10
    head_style = styles["BodyText"].clone("head", textColor=colors.white, fontName="Helvetica-Bold", fontSize=8)
    data = [[Paragraph(escape(c), head_style) for c in columns]]
    data += [[Paragraph(escape(str(c))[:600], cell_style) for c in r] for r in rows]
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f766e")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f4f6")]),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d1d5db")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    story = [Paragraph(escape(title), styles["Title"])]
    meta = f"{escape(subtitle)} · " if subtitle else ""
    story += [Paragraph(f"{meta}Generated {datetime.now():%Y-%m-%d %H:%M} · {len(rows)} row(s)", styles["Normal"]), Spacer(1, 10), table]
    doc.build(story)
    return out.getvalue(), FORMATS[fmt], filename
